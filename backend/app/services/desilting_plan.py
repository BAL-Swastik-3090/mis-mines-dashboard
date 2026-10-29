"""The mine's day-wise de-silting plan: read it, parse a workbook, store it.

WHY THIS EXISTS. De-silting was the one activity on the dashboard with a
permanent blank where every other had a target. Its actual comes from IMOS
(mines_day_wise_excavation, Variant '4', in CuM) and no plan for it existed
anywhere in any system, so the Executive Summary printed NO PLAN, the MIS KPI
card passed today_plan=None, and get_daily_actuals returned "silt_plan": None
with a comment saying why. The mine now keeps the plan in a spreadsheet, so
this takes that spreadsheet.

── THE PLAN IS IN POSTGRES AND THE ACTUAL IS IN MYSQL ──────────────────────
This table lives in minehub on Postgres, because it is the mine's own data
entered through this application. The actual it will be compared against lives
in balcorpdb on MySQL. THE TWO CANNOT BE JOINED IN SQL. Every caller therefore
reads the plan as a plain {date: qty} map and does the matching in Python —
which is why read_range() returns a dict rather than rows, and why no query in
this module mentions an actual.

── NOTHING IS WRITTEN UNTIL SOMEBODY HAS LOOKED AT IT ──────────────────────
parse() never touches the database. It reads a workbook and reports what it
found, including everything wrong with it, and the screen shows that back
before a single row is stored. diff() then says what confirming would change —
added, changed, removed — against what is already there. A plan people revise
mid-month is a plan somebody will one day re-upload from the wrong file, and
the moment to catch that is before the write, not in a review meeting.

── A MONTH IS REPLACED, NOT MERGED ─────────────────────────────────────────
On confirm the selected month becomes exactly what the file holds: dates in
the file are written, dates stored for that month and absent from the file are
deleted. Merging instead would leave a day the mine had deliberately removed
sitting on the dashboard with its old figure and nothing to show it was stale.
The diff names those removals explicitly so the confirmation is an informed
one. created_at and entry_id survive a re-upload; modified_at and modified_by
record the revision, which is what the mine asked for.
"""
from __future__ import annotations

import io
import logging
from calendar import monthrange
from datetime import date, datetime
from typing import Any

from sqlalchemy import text
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

#: How far down to look for the header. The mine's sheet starts at B3 — the
#: header is not at A1 and must not be assumed to be. Twenty rows is generous
#: enough for a title block and mean enough to fail on a sheet of something
#: else entirely.
_HEADER_SCAN_ROWS = 20

#: The word that ends the day rows. The mine's sheet carries a Total beneath
#: them, and a total read as a day would put a month's plan on one date.
_STOP_WORDS = {"total", "grand total", "sum", "month total"}

#: Quantities are CuM, the same unit as the actual, to three decimals — the
#: sheet divides a monthly target by the days in the month and arrives at
#: 645.1612903225806, so rounding to whole CuM would lose the reconciliation
#: with its own Total row.
_QTY_DP = 3

#: A Total row that disagrees with its own days by more than this is reported.
#: Not refused: the mine's sheet foots to 20,000.000000000004 against days
#: summing to 20,000.000000000001, and a rule strict enough to call that a
#: mismatch would reject every correct sheet it is ever given.
_TOTAL_TOLERANCE = 0.5


class SheetError(ValueError):
    """The workbook could not be read as a de-silting plan.

    Carries every problem found rather than the first, so somebody fixing a
    spreadsheet fixes it once instead of discovering the faults one upload at
    a time.
    """

    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("; ".join(problems))


# ── reading ─────────────────────────────────────────────────────────────────
def read_range(db: Session, from_date: date, to_date: date) -> dict[date, float]:
    """The plan for a date range, as {date: CuM}.

    A DATE WITH NO ROW IS ABSENT FROM THE MAP, never present as zero. Callers
    distinguish "the mine planned nothing that day" from "no plan has been
    uploaded for that month", and a defaulted zero would collapse the two —
    which is precisely the bug that makes a 0% achievement look like a failure
    rather than a missing upload.
    """
    rows = db.execute(text("""
        SELECT plan_date, plan_qty
        FROM   minehub.desilting_plan
        WHERE  plan_date BETWEEN :f AND :t
    """), {"f": from_date, "t": to_date}).all()
    return {r[0]: float(r[1]) for r in rows}


def plan_for(from_date: date, to_date: date) -> dict[date, float]:
    """read_range on a session of its own, and never raising.

    WHY NOT A SECOND DEPENDENCY ON EVERY ROUTE. The plan is in Postgres and
    every screen that wants it is built on MySQL, so each caller would need a
    second session threaded down from its router -- three routers, and every
    test that constructs those services by hand.

    WHY IT SWALLOWS THE ERROR. Postgres here is reached over the VPN and is
    not always up: this module was written on an afternoon when pg_hba.conf
    had no entry for the developer's new address. If that made the Executive
    Summary answer 500, a cold database would take down a page that is
    otherwise entirely MySQL. An unreachable plan is reported as NO PLAN --
    which is exactly what the screen showed before this feature existed, and
    is the truthful answer to "what is the target" when the table holding it
    cannot be read.
    """
    try:
        from app.minehub_db import SessionLocal
        if SessionLocal is None:
            return {}
        pg = SessionLocal()
    except Exception:                                   # noqa: BLE001
        return {}
    try:
        return read_range(pg, from_date, to_date)
    except Exception:                                   # noqa: BLE001
        logger.info("de-silting plan unavailable; reporting no plan",
                    exc_info=False)
        return {}
    finally:
        try:
            pg.close()
        except Exception:                               # noqa: BLE001
            pass


def read_month(db: Session, year: int, month: int) -> list[dict[str, Any]]:
    """Every stored row for one month, with who entered it and when."""
    first = date(year, month, 1)
    last = date(year, month, monthrange(year, month)[1])
    rows = db.execute(text("""
        SELECT plan_date, plan_qty, entry_id, created_at, modified_by, modified_at
        FROM   minehub.desilting_plan
        WHERE  plan_date BETWEEN :f AND :t
        ORDER BY plan_date
    """), {"f": first, "t": last}).mappings().all()
    return [{
        "plan_date":   r["plan_date"].isoformat(),
        "plan_qty":    float(r["plan_qty"]),
        "entry_id":    r["entry_id"],
        "created_at":  r["created_at"].isoformat() if r["created_at"] else None,
        "modified_by": r["modified_by"],
        "modified_at": r["modified_at"].isoformat() if r["modified_at"] else None,
    } for r in rows]


# ── parsing ─────────────────────────────────────────────────────────────────
def _as_date(v: Any) -> date | None:
    """A cell read as a date, or None if it is not one.

    Excel hands back a datetime for a date-formatted cell. A hand-typed date
    arrives as text, and is accepted in the two orders a person at this mine
    would actually type — day first, and ISO — but never in month-first order:
    03/04/2026 is ambiguous and guessing it wrong moves a plan by a month with
    nothing on screen to show it happened.
    """
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def _as_qty(v: Any) -> float | None:
    if isinstance(v, bool):           # bool is an int in Python; not a quantity
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _find_header(ws) -> tuple[int, int, int]:
    """(header row, date column, quantity column), all 1-based.

    Found by looking for the word "Date" rather than assumed to be the first
    row and the first column: the mine's sheet has its header at B3, with two
    empty rows and an empty column before it. Hard-coding A1 would read the
    header as a day and fail on the real file.
    """
    for r in range(1, min(ws.max_row, _HEADER_SCAN_ROWS) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().lower() == "date":
                # The quantity is the next column across that carries a
                # heading. Named rather than positional, so a spare column
                # between the two does not silently become the figure.
                for c2 in range(c + 1, ws.max_column + 1):
                    h = ws.cell(row=r, column=c2).value
                    if isinstance(h, str) and h.strip():
                        return r, c, c2
                raise SheetError([
                    'The sheet has a "Date" column but no quantity column '
                    "beside it. Expected a second heading such as "
                    '"De-Silting" to its right.'])
    raise SheetError([
        'No "Date" column was found in the first '
        f"{_HEADER_SCAN_ROWS} rows. Expected a header row with "
        '"Date" and "De-Silting".'])


def parse(payload: bytes, year: int, month: int) -> dict[str, Any]:
    """Read a workbook as one month's day-wise plan. Touches no database.

    THE MONTH IS THE USER'S, NOT THE FILE'S. The screen asks which month is
    being uploaded before the file is chosen, and a date outside it is an
    error rather than something to be accommodated. A sheet quietly carrying
    September while October was selected is the failure this exists to catch —
    it would otherwise overwrite a month nobody was looking at.
    """
    import openpyxl                                   # heavy; imported on use

    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    except Exception as e:                            # noqa: BLE001
        raise SheetError([f"That file could not be opened as an Excel "
                          f"workbook ({type(e).__name__})."]) from e

    ws = wb.worksheets[0]
    hdr, c_date, c_qty = _find_header(ws)
    qty_heading = str(ws.cell(row=hdr, column=c_qty).value or "").strip()

    first = date(year, month, 1)
    last = date(year, month, monthrange(year, month)[1])

    rows: list[dict[str, Any]] = []
    seen: dict[date, int] = {}
    problems: list[str] = []
    sheet_total: float | None = None

    for r in range(hdr + 1, ws.max_row + 1):
        raw_d = ws.cell(row=r, column=c_date).value
        raw_q = ws.cell(row=r, column=c_qty).value

        if raw_d is None and raw_q is None:
            continue

        # The Total row: keep its figure to check the days against, and stop.
        if isinstance(raw_d, str) and raw_d.strip().lower() in _STOP_WORDS:
            sheet_total = _as_qty(raw_q)
            break

        d = _as_date(raw_d)
        if d is None:
            problems.append(f"Row {r}: {raw_d!r} is not a date.")
            continue

        q = _as_qty(raw_q)
        if q is None:
            problems.append(
                f"Row {r} ({d:%d-%m-%Y}): {raw_q!r} is not a number.")
            continue
        if q < 0:
            problems.append(
                f"Row {r} ({d:%d-%m-%Y}): {q} is negative. A plan cannot be "
                "less than nothing.")
            continue

        if not (first <= d <= last):
            problems.append(
                f"Row {r}: {d:%d-%m-%Y} is not in {first:%B %Y}, the month "
                "selected for this upload.")
            continue
        if d in seen:
            problems.append(
                f"Row {r} ({d:%d-%m-%Y}): this date is already on row "
                f"{seen[d]}. A day can have only one plan.")
            continue

        seen[d] = r
        rows.append({"plan_date": d, "plan_qty": round(q, _QTY_DP)})

    if problems:
        raise SheetError(problems)
    if not rows:
        raise SheetError([
            f"No plan rows were found beneath the header on row {hdr}."])

    rows.sort(key=lambda x: x["plan_date"])
    total = round(sum(r["plan_qty"] for r in rows), _QTY_DP)

    # A Total that disagrees with its own days is REPORTED, NOT REFUSED. The
    # sheet is the mine's working document and its total is a convenience, not
    # the figure this stores; what gets stored is the days. Saying so lets
    # somebody notice they pasted over a row without being blocked by it.
    warnings: list[str] = []
    if sheet_total is not None and abs(sheet_total - total) > _TOTAL_TOLERANCE:
        warnings.append(
            f"The sheet's Total row says {sheet_total:,.2f} but its days add "
            f"up to {total:,.2f}. The days are what will be stored.")

    missing = [d for d in (first.replace(day=i + 1)
                           for i in range(monthrange(year, month)[1]))
               if d not in seen]
    if missing:
        warnings.append(
            f"{len(missing)} day(s) of {first:%B} are not in the sheet "
            f"({', '.join(f'{d:%d}' for d in missing[:8])}"
            f"{', …' if len(missing) > 8 else ''}). They will carry no plan.")

    return {
        "year": year, "month": month,
        "qty_heading": qty_heading,
        "header_row": hdr,
        "rows": rows,
        "total": total,
        "sheet_total": sheet_total,
        "warnings": warnings,
    }


# ── what confirming would change ────────────────────────────────────────────
def diff(db: Session, parsed: dict[str, Any]) -> dict[str, Any]:
    """Added, changed and removed, against what is stored for that month.

    Computed BEFORE the write and shown on screen, because "replace the month"
    is only a safe rule if the person clicking can see what it replaces.
    """
    existing = read_month(db, parsed["year"], parsed["month"])
    have = {date.fromisoformat(r["plan_date"]): r["plan_qty"] for r in existing}
    incoming = {r["plan_date"]: r["plan_qty"] for r in parsed["rows"]}

    added, changed, unchanged = [], [], 0
    for d, q in sorted(incoming.items()):
        if d not in have:
            added.append({"plan_date": d.isoformat(), "plan_qty": q})
        elif abs(have[d] - q) > 10 ** -_QTY_DP:
            changed.append({"plan_date": d.isoformat(),
                            "was": have[d], "now": q})
        else:
            unchanged += 1
    removed = [{"plan_date": d.isoformat(), "was": q}
               for d, q in sorted(have.items()) if d not in incoming]

    return {
        "added": added, "changed": changed, "removed": removed,
        "unchanged": unchanged,
        "had_rows": len(have),
        "is_revision": bool(have),
        "total_was": round(sum(have.values()), _QTY_DP) if have else None,
        "total_now": parsed["total"],
    }


# ── writing ─────────────────────────────────────────────────────────────────
def replace_month(db: Session, parsed: dict[str, Any], actor: str) -> dict[str, Any]:
    """Make the month exactly what the file holds. One transaction.

    ONE TRANSACTION, because a half-written month is worse than an unwritten
    one: the dashboard would report a plan for the first fortnight and nothing
    for the second, and the figure would look like a mine that stopped
    de-silting rather than an upload that stopped halfway.

    created_at and entry_id are NOT touched on a row that already exists --
    they record who first entered the plan. modified_at and modified_by carry
    the revision, which is what the mine asked to be able to see.
    """
    year, month = parsed["year"], parsed["month"]
    first = date(year, month, 1)
    last = date(year, month, monthrange(year, month)[1])
    changes = diff(db, parsed)

    keep = [r["plan_date"] for r in parsed["rows"]]
    # Dates the mine has dropped from the plan. Deleted rather than zeroed: a
    # zero is a plan of nothing, and that is a different statement.
    db.execute(text("""
        DELETE FROM minehub.desilting_plan
        WHERE plan_date BETWEEN :f AND :t
          AND NOT (plan_date = ANY(:keep))
    """), {"f": first, "t": last, "keep": keep})

    for r in parsed["rows"]:
        # ALIASED, because the WHERE of an ON CONFLICT DO UPDATE must name the
        # target by its bare table name or an alias -- a schema-qualified
        # minehub.desilting_plan.plan_qty there is not a column reference
        # Postgres will resolve. `dp` makes it unambiguous either way.
        #
        # The WHERE is what keeps modified_at honest: a day whose figure is
        # unchanged is not written at all, so re-uploading the same sheet twice
        # does not stamp thirty-one spurious modifications.
        db.execute(text("""
            INSERT INTO minehub.desilting_plan AS dp
                   (plan_date, plan_qty, entry_id, created_at)
            VALUES (:d, :q, :by, now())
            ON CONFLICT (plan_date) DO UPDATE
               SET plan_qty    = EXCLUDED.plan_qty,
                   modified_by = EXCLUDED.entry_id,
                   modified_at = now()
             WHERE dp.plan_qty IS DISTINCT FROM EXCLUDED.plan_qty
        """), {"d": r["plan_date"], "q": r["plan_qty"], "by": actor})

    db.commit()
    logger.info("de-silting plan for %04d-%02d written by %s: "
                "%d added, %d changed, %d removed",
                year, month, actor, len(changes["added"]),
                len(changes["changed"]), len(changes["removed"]))
    return changes
